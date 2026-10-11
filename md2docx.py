#!/usr/bin/env python3
"""
md2docx — Convert Markdown to a clean, well-formatted Word (.docx) document.

Built for Markdown written by people and by AI assistants (ChatGPT, etc.),
with proper support for Chinese / CJK text.

Usage:
    python md2docx.py input.md                    -> input.docx
    python md2docx.py input.md output.docx        -> output.docx
    python md2docx.py input.md -o report.docx     -> report.docx
    python md2docx.py a.md b.md "notes/*.md"      -> one .docx per file
    python md2docx.py input.md --toc --cjk-font 宋体

Library use:
    from md2docx import md_to_docx, convert
    md_to_docx("input.md", "output.docx")
    doc = convert("# Hello\\n\\n你好，世界")   # returns a python-docx Document

Requires: python-docx >= 1.0   (pip install python-docx)
"""

from __future__ import annotations

import argparse
import glob
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Optional
from urllib.parse import unquote

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor, Twips

__version__ = "2.0.0"

# ─────────────────────────────────────────────
# Defaults
# ─────────────────────────────────────────────

DEFAULT_FONT = "Calibri"
DEFAULT_CJK_FONT = "等线"          # DengXian — Office's default CJK pairing for Calibri
DEFAULT_CODE_FONT = "Consolas"

HEADING_COLOR = RGBColor(0x1F, 0x38, 0x64)
CODE_COLOR = RGBColor(0xC0, 0x39, 0x2B)
LINK_COLOR = RGBColor(0x05, 0x63, 0xC1)
MUTED_COLOR = RGBColor(0x59, 0x59, 0x59)
PLACEHOLDER_COLOR = RGBColor(0x88, 0x88, 0x88)
TABLE_HEADER_FILL = "D9E2F3"

# Heading size relative to body size (11pt body -> ~20 / 16 / 13 / 11 / 11 / 11)
HEADING_SCALE = {1: 1.8, 2: 1.45, 3: 1.2, 4: 1.0, 5: 1.0, 6: 1.0}

# List geometry (twips: 1 cm = 567 twips)
LIST_LEFT = 720      # text position of level 0
LIST_STEP = 432      # extra indent per nesting level
LIST_HANG = 360      # space between marker and text
BULLET_GLYPHS = [("•", "Arial"), ("o", "Courier New"), ("▪", "Arial")]
NUMBER_FORMATS = ["decimal", "lowerLetter", "lowerRoman"]

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}

ALERT_LABELS = {
    "NOTE": "Note", "TIP": "Tip", "IMPORTANT": "Important",
    "WARNING": "Warning", "CAUTION": "Caution",
}


@dataclass
class Options:
    """Conversion options. Fonts left as None fall back to the defaults
    (or, when a template is given, to whatever the template defines)."""
    font: Optional[str] = None
    cjk_font: Optional[str] = None
    code_font: str = DEFAULT_CODE_FONT
    font_size: float = 11
    toc: bool = False
    toc_title: str = "Contents"
    hr_pagebreak: bool = False
    template: Optional[str] = None


# ─────────────────────────────────────────────
# Block-level patterns
# ─────────────────────────────────────────────

FENCE_RE = re.compile(r"^(\s*)(`{3,}|~{3,})(.*)$")
HR_RE = re.compile(r"^ {0,3}([-*_])(?:[ \t]*\1){2,}[ \t]*$")
HEADING_RE = re.compile(r"^ {0,3}(#{1,6})(?:[ \t]+(.*?))?(?:[ \t]+#+)?[ \t]*$")
SETEXT_H1_RE = re.compile(r"^ {0,3}=+[ \t]*$")
QUOTE_RE = re.compile(r"^ {0,3}>")
BULLET_RE = re.compile(r"^(\s*)([-*+])[ \t]+(.*)$")
ORDERED_RE = re.compile(r"^(\s*)(\d{1,9})([.)])[ \t]+(.*)$")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-+:?\s*(?:\|\s*:?-+:?\s*)*\|?\s*$")
IMAGE_LINE_RE = re.compile(r'^!\[([^\]]*)\]\(\s*<?([^)\s>]+)>?(?:\s+"[^"]*")?\s*\)$')
# Author convention: a line like [**Figure 1: caption**] becomes a centred placeholder
PLACEHOLDER_RE = re.compile(r"^\[(\*{1,3}[^\[\]]+?\*{1,3})\]$")
TASK_RE = re.compile(r"^\[([ xX])\][ \t]+")
ALERT_RE = re.compile(r"^\[!(NOTE|TIP|IMPORTANT|WARNING|CAUTION)\]\s*$", re.I)

# ─────────────────────────────────────────────
# Inline patterns (leftmost match wins; at the same position, earlier wins)
# ─────────────────────────────────────────────

_INLINE_PARTS = [
    r"(?P<code>(?P<ticks>`+)(?P<code_text>.+?)(?<!`)(?P=ticks)(?!`))",
    r"(?P<escape>\\(?P<esc>[\\`*_{}\[\]()#+\-.!|~<>]))",
    r"(?P<br><br\s*/?>)",
    r'(?P<image>!\[(?P<img_alt>[^\]]*)\]\(\s*<?(?P<img_src>[^)\s>]+)>?(?:\s+"[^"]*")?\s*\))',
    r'(?P<link>\[(?P<link_text>(?:\\.|[^\[\]\\]|\[[^\]]*\])+)\]\(\s*<?(?P<link_url>[^)\s>]+)>?(?:\s+"[^"]*")?\s*\))',
    r"(?P<autolink><(?P<auto_url>(?:https?|ftp)://[^>\s]+|mailto:[^>\s]+)>)",
    r"(?P<url>(?<![\w/@.])(?:https?://|www\.)[^\s<>\"'\[\]()（）]*"
    r"[^\s<>\"'\[\]()（）.,;:!?，。；：！？、》」』…])",
    r"(?P<strong_em>\*\*\*(?=\S)(?P<se_text>.+?)(?<=\S)\*\*\*)",
    r"(?P<strong>\*\*(?=\S)(?P<s_text>.+?)(?<=\S)\*\*"
    r"|(?<![\w\\])__(?=\S)(?P<s_text2>.+?)(?<=\S)__(?!\w))",
    r"(?P<em>\*(?=[^\s*])(?P<e_text>.+?)(?<=[^\s*\\])\*"
    r"|(?<![\w\\])_(?=\S)(?P<e_text2>.+?)(?<=\S)_(?!\w))",
    r"(?P<strike>~~(?=\S)(?P<st_text>.+?)(?<=\S)~~)",
]
INLINE_RE = re.compile("|".join(_INLINE_PARTS))

# Chinese / Japanese characters and full-width punctuation. Lines are joined
# without a space between two such characters (Korean uses spaces, so Hangul
# is deliberately not included).
_CJK_RE = re.compile(
    "[⺀-⿟　-〿぀-ヿ㄀-ㄯ㆐-ㇿ"
    "㐀-䶿一-鿿豈-﫿︰-﹏＀-￯]"
)


# ─────────────────────────────────────────────
# Low-level XML helpers
# ─────────────────────────────────────────────

# Child order required by the OOXML schema. Word rejects files whose
# elements are out of order, so every manual insert goes through these.
_PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
    "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
    "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
    "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]
_RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
    "dstrike", "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid",
    "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs",
    "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs",
    "em", "lang", "eastAsianLayout", "specVanish", "oMath",
]
_TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
    "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
]
# Settings elements that must come *after* <w:updateFields>
_SETTINGS_AFTER_UPDATEFIELDS = {
    "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
    "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
    "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
    "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
    "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator",
}


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _insert_ordered(parent, child, order):
    """Insert `child` into `parent`, replacing any same-named element and
    keeping schema order."""
    name = _local(child.tag)
    for existing in parent.findall(child.tag):
        parent.remove(existing)
    later = set(order[order.index(name) + 1:])
    for existing in parent:
        if _local(existing.tag) in later:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


def _set_fonts(rPr, latin: Optional[str], cjk: Optional[str]):
    """Set explicit Latin and East Asian fonts on an rPr element.

    python-docx's `font.name` only sets w:ascii / w:hAnsi. Chinese text uses
    w:eastAsia, and any *Theme attribute overrides the explicit font, so the
    theme attributes have to be removed as well."""
    if rPr is None or not (latin or cjk):
        return
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = _insert_ordered(rPr, OxmlElement("w:rFonts"), _RPR_ORDER)
    if latin:
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme"):
            rFonts.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rFonts.set(qn(attr), latin)
    if cjk:
        rFonts.attrib.pop(qn("w:eastAsiaTheme"), None)
        rFonts.set(qn("w:eastAsia"), cjk)


def _border(sides, sz=4, space=4, color="D0D7DE"):
    inner = "".join(
        f'<w:{s} w:val="single" w:sz="{sz}" w:space="{space}" w:color="{color}"/>'
        for s in sides
    )
    return parse_xml(f"<w:pBdr {nsdecls('w')}>{inner}</w:pBdr>")


def _shading(fill):
    return parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')


def _indent_width(line: str) -> int:
    expanded = line.expandtabs(4)
    return len(expanded) - len(expanded.lstrip(" "))


def _is_cjk(ch: str) -> bool:
    return bool(_CJK_RE.match(ch))


def _join_lines(lines):
    """Join the raw lines of one paragraph.

    - A line ending in two spaces or a backslash forces a line break.
    - Two lines that meet between CJK characters are joined with no space
      (otherwise Chinese text gets stray spaces wherever the source wrapped).
    """
    out = ""
    hard_break = False
    for k, raw in enumerate(lines):
        raw = raw.rstrip("\r")
        seg = raw.strip()
        this_hard = raw.endswith("  ") or (seg.endswith("\\") and not seg.endswith("\\\\"))
        if seg.endswith("\\") and not seg.endswith("\\\\"):
            seg = seg[:-1].rstrip()
        if k == 0:
            out = seg
        elif hard_break:
            out += "\n" + seg
        elif out and seg and _is_cjk(out[-1]) and _is_cjk(seg[0]):
            out += seg
        elif out and seg:
            out += " " + seg
        else:
            out += seg
        hard_break = this_hard
    return out


def _split_table_row(line: str):
    """Split a table row on unescaped pipes that are not inside `code`."""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    cells, cur, j = [], [], 0
    while j < len(s):
        c = s[j]
        if c == "\\" and j + 1 < len(s):
            cur.append("|" if s[j + 1] == "|" else s[j:j + 2])
            j += 2
            continue
        if c == "`":
            code = INLINE_RE.match(s, j)
            if code is not None and code.group("code") is not None:
                # Use the same code-span rules as inline rendering, including
                # matching multi-backtick delimiters and literal backticks.
                cur.append(code.group("code").replace("\\|", "|"))
                j = code.end()
                continue
        if c == "|":
            cells.append("".join(cur).strip())
            cur = []
        else:
            cur.append(c)
        j += 1
    cells.append("".join(cur).strip())
    return cells


# ─────────────────────────────────────────────
# Converter
# ─────────────────────────────────────────────

class Converter:
    def __init__(self, options: Optional[Options] = None, base_dir: str | Path = "."):
        self.o = options or Options()
        self.base_dir = Path(base_dir)
        using_template = bool(self.o.template)
        # With a template, fonts are only overridden when explicitly requested.
        self.font = self.o.font or (None if using_template else DEFAULT_FONT)
        self.cjk_font = self.o.cjk_font or (None if using_template else DEFAULT_CJK_FONT)
        self.doc = self._new_document()
        self._abstract_ids = {}
        self._bullet_num_id = None
        self.list_style = self._style_or_none("List Paragraph")

    # ── document & styles ──

    def _new_document(self):
        if self.o.template:
            doc = Document(self.o.template)
            body = doc.element.body
            for child in list(body):
                if child.tag != qn("w:sectPr"):
                    body.remove(child)
        else:
            doc = Document()
            for section in doc.sections:
                section.top_margin = section.bottom_margin = Cm(2.54)
                section.left_margin = section.right_margin = Cm(2.54)
        self.doc = doc
        self._setup_styles(restyle=not self.o.template)
        return doc

    def _style_or_none(self, name):
        try:
            return self.doc.styles[name]
        except KeyError:
            return None

    def _setup_styles(self, restyle: bool):
        doc, o = self.doc, self.o
        styles = doc.styles

        # Document-wide defaults (theme fonts would otherwise win)
        defaults = doc.styles.element.find(qn("w:docDefaults"))
        if defaults is not None:
            rpr = defaults.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}")
            _set_fonts(rpr, self.font, self.cjk_font)
        if self.cjk_font:
            # python-docx's template declares ja-JP, which makes Word pick a
            # Japanese theme font for CJK text in some styles.
            lang = doc.settings.element.find(qn("w:themeFontLang"))
            if lang is not None:
                lang.set(qn("w:eastAsia"), "zh-CN")
        zoom = doc.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")  # required by the schema

        normal = styles["Normal"]
        _set_fonts(normal.element.get_or_add_rPr(), self.font, self.cjk_font)
        if restyle:
            normal.font.size = Pt(o.font_size)
            normal.paragraph_format.space_after = Pt(6)
            normal.paragraph_format.line_spacing = 1.15

        for level, scale in HEADING_SCALE.items():
            s = self._style_or_none(f"Heading {level}")
            if s is None:
                continue
            _set_fonts(s.element.get_or_add_rPr(), self.font, self.cjk_font)
            if restyle:
                s.font.size = Pt(round(o.font_size * scale * 2) / 2)
                s.font.bold = level <= 5
                s.font.italic = False
                s.font.color.rgb = HEADING_COLOR
                s.paragraph_format.space_before = Pt(12)
                s.paragraph_format.space_after = Pt(4)

        # Custom styles — created only if missing, so a template can define its own.
        quote, created = self._ensure_style("Block Quote")
        if created:
            quote.font.color.rgb = MUTED_COLOR
            quote.paragraph_format.left_indent = Cm(0.6)
            quote.paragraph_format.space_after = Pt(6)
            _insert_ordered(quote.element.get_or_add_pPr(),
                            _border(["left"], sz=18, space=8), _PPR_ORDER)

        code, created = self._ensure_style("Code Block")
        if created:
            code.font.size = Pt(max(o.font_size - 1.5, 6))
            code.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
            pf = code.paragraph_format
            pf.space_before, pf.space_after = Pt(4), Pt(8)
            pf.line_spacing = 1.0
            pf.left_indent = pf.right_indent = Cm(0.2)
            ppr = code.element.get_or_add_pPr()
            _insert_ordered(ppr, _border(["top", "left", "bottom", "right"], color="E1E4E8"),
                            _PPR_ORDER)
            _insert_ordered(ppr, _shading("F6F8FA"), _PPR_ORDER)
        _set_fonts(code.element.get_or_add_rPr(), o.code_font, self.cjk_font)

    def _ensure_style(self, name):
        existing = self._style_or_none(name)
        if existing is not None:
            return existing, False
        style = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = self.doc.styles["Normal"]
        style.quick_style = True
        return style, True

    def _separate_from(self, style_name):
        """Word merges the borders of consecutive paragraphs that share a
        style, so two quotes (or code blocks) in a row would look like one.
        Insert a tiny spacer paragraph when the previous block uses the same style."""
        style = self._style_or_none(style_name)
        body = self.doc.element.body
        last = None
        for child in reversed(body):
            if child.tag != qn("w:sectPr"):
                last = child
                break
        if style is None or last is None or last.tag != qn("w:p"):
            return
        pStyle = last.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
        if pStyle is not None and pStyle.get(qn("w:val")) == style.style_id:
            spacer = self.doc.add_paragraph()
            spacer.paragraph_format.space_before = Pt(0)
            spacer.paragraph_format.space_after = Pt(0)
            spacer.paragraph_format.line_spacing = Pt(2)

    def _text_width(self):
        sec = self.doc.sections[-1]
        try:
            return sec.page_width - sec.left_margin - sec.right_margin
        except TypeError:
            return Cm(16)

    # ── numbering (real Word lists, with nesting and restarts) ──

    def _numbering(self):
        return self.doc.part.numbering_part.element

    def _abstract_num(self, kind):
        if kind in self._abstract_ids:
            return self._abstract_ids[kind]
        numbering = self._numbering()
        ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
        new_id = max(ids, default=-1) + 1
        levels = []
        for lvl in range(9):
            left = LIST_LEFT + LIST_STEP * lvl
            ind = f'<w:pPr><w:ind w:left="{left}" w:hanging="{LIST_HANG}"/></w:pPr>'
            if kind == "bullet":
                glyph, font = BULLET_GLYPHS[lvl % len(BULLET_GLYPHS)]
                levels.append(
                    f'<w:lvl w:ilvl="{lvl}"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
                    f'<w:lvlText w:val="{glyph}"/><w:lvlJc w:val="left"/>{ind}'
                    f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:hint="default"/></w:rPr>'
                    f"</w:lvl>"
                )
            else:
                fmt = NUMBER_FORMATS[lvl % len(NUMBER_FORMATS)]
                levels.append(
                    f'<w:lvl w:ilvl="{lvl}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
                    f'<w:lvlText w:val="%{lvl + 1}."/><w:lvlJc w:val="left"/>{ind}</w:lvl>'
                )
        el = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{new_id}">'
            f'<w:multiLevelType w:val="multilevel"/>{"".join(levels)}</w:abstractNum>'
        )
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(el)
        else:
            numbering.append(el)
        self._abstract_ids[kind] = new_id
        return new_id

    def _new_num(self, kind, level=0, start=None):
        numbering = self._numbering()
        abstract_id = self._abstract_num(kind)
        ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        num_id = max(ids, default=0) + 1
        override = ""
        if start is not None:
            override = (f'<w:lvlOverride w:ilvl="{level}">'
                        f'<w:startOverride w:val="{start}"/></w:lvlOverride>')
        el = parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{num_id}">'
            f'<w:abstractNumId w:val="{abstract_id}"/>{override}</w:num>'
        )
        cleanup = numbering.find(qn("w:numIdMacAtCleanup"))
        if cleanup is not None:
            cleanup.addprevious(el)
        else:
            numbering.append(el)
        return num_id

    def _bullet_num(self):
        if self._bullet_num_id is None:
            self._bullet_num_id = self._new_num("bullet")
        return self._bullet_num_id

    def _list_paragraph(self, num_id, level):
        p = self.doc.add_paragraph(style=self.list_style)
        numPr = parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level}"/>'
            f'<w:numId w:val="{num_id}"/></w:numPr>'
        )
        _insert_ordered(p._p.get_or_add_pPr(), numPr, _PPR_ORDER)
        return p

    # ── main loop ──

    def convert(self, md_text: str):
        lines = md_text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        i = self._skip_front_matter(lines)
        if self.o.toc:
            self._add_toc()
        n = len(lines)
        while i < n:
            line = lines[i]
            s = line.strip()
            if not s:
                i += 1
                continue
            if s.startswith("<!--"):
                i = self._skip_html_comment(lines, i)
                continue
            m = FENCE_RE.match(line)
            if m:
                i, _ = self._code_block(lines, i, m)
                continue
            if HR_RE.match(line):
                self._hr()
                i += 1
                continue
            m = HEADING_RE.match(line)
            if m:
                self._heading(len(m.group(1)), m.group(2) or "")
                i += 1
                continue
            if QUOTE_RE.match(line):
                i = self._blockquote(lines, i)
                continue
            if self._is_table_start(lines, i):
                i = self._table(lines, i)
                continue
            if self._is_list_item(line):
                i = self._list(lines, i)
                continue
            m = IMAGE_LINE_RE.match(s)
            if m:
                self._block_image(m.group(1), m.group(2))
                i += 1
                continue
            m = PLACEHOLDER_RE.match(s)
            if m:
                self._placeholder(m.group(1))
                i += 1
                continue
            i = self._paragraph(lines, i)
        return self.doc

    # ── block detection ──

    @staticmethod
    def _is_list_item(line):
        return bool((BULLET_RE.match(line) or ORDERED_RE.match(line)) and not HR_RE.match(line))

    @staticmethod
    def _is_table_start(lines, i):
        if i + 1 >= len(lines):
            return False
        head, sep = lines[i], lines[i + 1]
        return "|" in head and "|" in sep and "-" in sep and bool(TABLE_SEP_RE.match(sep))

    def _is_block_start(self, lines, i):
        line = lines[i]
        s = line.strip()
        return bool(
            FENCE_RE.match(line) or HR_RE.match(line) or HEADING_RE.match(line)
            or QUOTE_RE.match(line) or self._is_list_item(line)
            or self._is_table_start(lines, i) or IMAGE_LINE_RE.match(s)
            or PLACEHOLDER_RE.match(s) or s.startswith("<!--")
        )

    @staticmethod
    def _skip_front_matter(lines):
        """Skip a YAML front-matter block (--- key: value ... ---) at the top."""
        if not lines or lines[0].strip() != "---" or len(lines) < 3:
            return 0
        if not re.match(r"^[\w-]+\s*:", lines[1]):
            return 0
        for j in range(1, min(len(lines), 200)):
            if lines[j].strip() in ("---", "..."):
                return j + 1
        return 0

    @staticmethod
    def _skip_html_comment(lines, i):
        while i < len(lines):
            if "-->" in lines[i]:
                return i + 1
            i += 1
        return i

    # ── blocks ──

    def _heading(self, level, text):
        try:
            p = self.doc.add_heading(level=level)
            self._inline(p, text, code_shrink=False)
        except KeyError:  # template without heading styles
            p = self.doc.add_paragraph()
            self._inline(p, text, bold=True, code_shrink=False)

    def _paragraph(self, lines, i):
        buf = [lines[i]]
        i += 1
        while i < len(lines):
            line = lines[i]
            if not line.strip():
                break
            if SETEXT_H1_RE.match(line):          # Title\n=====
                self._heading(1, _join_lines(buf))
                return i + 1
            if self._is_block_start(lines, i):
                break
            buf.append(line)
            i += 1
        p = self.doc.add_paragraph()
        self._inline(p, _join_lines(buf))
        return i

    def _hr(self):
        if self.o.hr_pagebreak:
            self.doc.add_page_break()
            return
        p = self.doc.add_paragraph()
        _insert_ordered(p._p.get_or_add_pPr(),
                        _border(["bottom"], sz=6, space=1, color="BFBFBF"), _PPR_ORDER)

    def _code_block(self, lines, i, m):
        indent = _indent_width(m.group(1))
        fence = m.group(2)
        i += 1
        code = []
        while i < len(lines):
            line = lines[i]
            close = re.match(r"^\s*(`{3,}|~{3,})\s*$", line)
            if close and close.group(1)[0] == fence[0] and len(close.group(1)) >= len(fence):
                i += 1
                break
            # remove the fence's own indentation (e.g. code inside a list item)
            strip = min(indent, _indent_width(line))
            code.append(line.expandtabs(4)[strip:] if strip else line)
            i += 1
        while code and not code[-1].strip():
            code.pop()
        self._separate_from("Code Block")
        p = self.doc.add_paragraph(style="Code Block")
        if code:
            p.add_run("\n".join(code))  # python-docx turns \n into line breaks
        return i, p

    def _blockquote(self, lines, i):
        content = []
        while i < len(lines):
            line = lines[i]
            if QUOTE_RE.match(line):
                content.append(re.sub(r"^(?:\s{0,3}>\s?)+", "", line))
                i += 1
            elif line.strip() and content and content[-1].strip() \
                    and not self._is_block_start(lines, i):
                content.append(line)  # lazy continuation
                i += 1
            else:
                break

        # GitHub alert syntax: > [!NOTE]
        label = None
        if content and ALERT_RE.match(content[0].strip()):
            label = ALERT_LABELS[ALERT_RE.match(content[0].strip()).group(1).upper()]
            content = content[1:]

        groups, cur = [], []
        for line in content:
            if line.strip():
                cur.append(line)
            elif cur:
                groups.append(cur)
                cur = []
        if cur:
            groups.append(cur)

        self._separate_from("Block Quote")
        if label:
            p = self.doc.add_paragraph(style="Block Quote")
            p.add_run(label).bold = True
        for group in groups:
            # keep list-like lines on their own line inside the quote
            text = ""
            for k, line in enumerate(group):
                lm = BULLET_RE.match(line)
                piece = ("• " + lm.group(3)) if lm and not HR_RE.match(line) else line
                if k == 0:
                    text = _join_lines([piece])
                elif lm or ORDERED_RE.match(line):
                    text += "\n" + piece.strip()
                else:
                    text = _join_lines([text, piece])
            p = self.doc.add_paragraph(style="Block Quote")
            self._inline(p, text)
        return i

    def _table(self, lines, i):
        header = _split_table_row(lines[i])
        aligns = []
        for cell in _split_table_row(lines[i + 1]):
            c = cell.strip()
            if c.startswith(":") and c.endswith(":"):
                aligns.append(WD_ALIGN_PARAGRAPH.CENTER)
            elif c.endswith(":"):
                aligns.append(WD_ALIGN_PARAGRAPH.RIGHT)
            else:
                aligns.append(None)
        i += 2
        rows = []
        while i < len(lines) and lines[i].strip() and "|" in lines[i]:
            rows.append(_split_table_row(lines[i]))
            i += 1

        ncols = len(header)
        table = self.doc.add_table(rows=1 + len(rows), cols=ncols)
        if self._style_or_none("Table Grid") is not None:
            table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_size = max(self.o.font_size - 1, 6)

        for r, row in enumerate([header] + rows):
            row = (row + [""] * ncols)[:ncols]
            for c, text in enumerate(row):
                cell = table.rows[r].cells[c]
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                if c < len(aligns) and aligns[c] is not None:
                    p.alignment = aligns[c]
                self._inline(p, text, size=cell_size, bold=(r == 0))
                if r == 0:
                    _insert_ordered(cell._tc.get_or_add_tcPr(), _shading(TABLE_HEADER_FILL),
                                    _TCPR_ORDER)
        # repeat header row on every page
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        self.doc.add_paragraph()  # keeps consecutive tables from merging
        return i

    def _list(self, lines, i):
        """Bullet / numbered lists, including nesting, mixed types, numbering
        that restarts per list (and respects '3.' as a start number),
        continuation paragraphs and code blocks inside list items."""
        n = len(lines)
        stack = []        # indentation of each open nesting level
        seq = {}          # level -> numId of the numbered sequence in progress
        while i < n:
            line = lines[i]

            if self._is_list_item(line):
                m = ORDERED_RE.match(line) or BULLET_RE.match(line)
                ordered = m.re is ORDERED_RE
                indent = _indent_width(m.group(1))
                while stack and indent < stack[-1]:
                    stack.pop()
                if not stack or indent > stack[-1]:
                    stack.append(indent)
                level = min(len(stack) - 1, 8)
                for deeper in [k for k in seq if k > level]:
                    del seq[deeper]

                if ordered:
                    text = m.group(4)
                    num_id = seq.get(level)
                    if num_id is None:
                        num_id = self._new_num("ordered", level, int(m.group(2)))
                        seq[level] = num_id
                else:
                    text = m.group(3)
                    num_id = self._bullet_num()
                    seq[level] = None

                buf = [text]
                i += 1
                while i < n and lines[i].strip() and not self._is_block_start(lines, i):
                    buf.append(lines[i])
                    i += 1
                text = _join_lines(buf)
                task = TASK_RE.match(text)
                if task:
                    text = ("☑ " if task.group(1) in "xX" else "☐ ") + text[task.end():]
                p = self._list_paragraph(num_id, level)
                self._inline(p, text)
                continue

            if not line.strip():
                j = i + 1
                while j < n and not lines[j].strip():
                    j += 1
                if j < n and (self._is_list_item(lines[j])
                              or (stack and _indent_width(lines[j]) > stack[0])):
                    i = j
                    continue
                break

            # Indented content that belongs to a list item
            ind = _indent_width(line)
            if stack and ind > stack[0]:
                level = max(k for k, s in enumerate(stack) if s < ind)
                left = Twips(LIST_LEFT + LIST_STEP * min(level, 8))
                fm = FENCE_RE.match(line)
                if fm:
                    i, p = self._code_block(lines, i, fm)
                    p.paragraph_format.left_indent = left
                    continue
                if self._is_block_start(lines, i):
                    break
                buf = [line]
                i += 1
                while i < n and lines[i].strip() and not self._is_block_start(lines, i):
                    buf.append(lines[i])
                    i += 1
                p = self.doc.add_paragraph(style=self.list_style)
                p.paragraph_format.left_indent = left
                self._inline(p, _join_lines(buf))
                continue
            break
        return i

    def _resolve_image(self, src):
        if re.match(r"^[a-z][a-z0-9+.-]*://", src, re.I):
            return None
        path = Path(unquote(src))
        if not path.is_absolute():
            path = self.base_dir / path
        if path.suffix.lower() not in IMAGE_EXTS or not path.is_file():
            return None
        return path

    def _add_picture(self, run, src):
        path = self._resolve_image(src)
        if path is None:
            return False
        try:
            shape = run.add_picture(str(path))
        except Exception as exc:  # unsupported/corrupt image
            print(f"Warning: could not embed image {path}: {exc}", file=sys.stderr)
            return False
        max_w = self._text_width()
        if shape.width > max_w:
            shape.height = int(shape.height * max_w / shape.width)
            shape.width = max_w
        return True

    def _image_fallback(self, p, alt, src):
        if not re.match(r"^[a-z][a-z0-9+.-]*://", src, re.I):
            print(f"Warning: image not found or unsupported: {src}", file=sys.stderr)
        label = f"[Image: {alt}]" if alt else "[Image]"
        self._emit(p, [("text", label, {"italic": True, "link": src})], muted=True)

    def _block_image(self, alt, src):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = bool(alt)
        if not self._add_picture(p.add_run(), src):
            self._image_fallback(p, alt, src)
            return
        if alt:
            cap = self.doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._inline(cap, alt, size=max(self.o.font_size - 1.5, 6), italic=True, muted=True)

    def _placeholder(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(12)
        self._inline(p, text, italic=True, color=PLACEHOLDER_COLOR)

    def _add_toc(self):
        if self._style_or_none("TOC Heading") is not None:
            title = self.doc.add_paragraph(self.o.toc_title, style="TOC Heading")
        else:
            title = self.doc.add_paragraph()
            run = title.add_run(self.o.toc_title)
            run.bold = True
            run.font.size = Pt(round(self.o.font_size * 1.45))
        title.paragraph_format.space_after = Pt(6)

        p = self.doc.add_paragraph()
        w = nsdecls("w")
        for xml in (
            f'<w:r {w}><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>',
            f'<w:r {w}><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>',
            f'<w:r {w}><w:fldChar w:fldCharType="separate"/></w:r>',
            f'<w:r {w}><w:t>Right-click here and choose "Update Field" to build the table of contents.</w:t></w:r>',
            f'<w:r {w}><w:fldChar w:fldCharType="end"/></w:r>',
        ):
            p._p.append(parse_xml(xml))

        # Ask Word to refresh fields when the file is opened
        settings = self.doc.settings.element
        if settings.find(qn("w:updateFields")) is None:
            el = parse_xml(f'<w:updateFields {w} w:val="true"/>')
            for child in settings:
                if _local(child.tag) in _SETTINGS_AFTER_UPDATEFIELDS:
                    child.addprevious(el)
                    break
            else:
                settings.append(el)

    # ── inline ──

    def _parse_inline(self, text, fmt):
        out, pos = [], 0
        for m in INLINE_RE.finditer(text):
            if m.start() > pos:
                out.append(("text", text[pos:m.start()], fmt))
            pos = m.end()
            g = m.group
            if g("code") is not None:
                code = g("code_text")
                if len(code) > 2 and code[0] == " " and code[-1] == " " and code.strip():
                    code = code[1:-1]
                out.append(("text", code, {**fmt, "code": True}))
            elif g("escape") is not None:
                out.append(("text", g("esc"), fmt))
            elif g("br") is not None:
                out.append(("text", "\n", fmt))
            elif g("image") is not None:
                out.append(("image", g("img_alt"), g("img_src"), fmt))
            elif g("link") is not None:
                out += self._parse_inline(g("link_text"), {**fmt, "link": g("link_url")})
            elif g("autolink") is not None:
                out.append(("text", g("auto_url"), {**fmt, "link": g("auto_url")}))
            elif g("url") is not None:
                url = g("url")
                target = url if "://" in url else "http://" + url
                out.append(("text", url, {**fmt, "link": target}))
            elif g("strong_em") is not None:
                out += self._parse_inline(g("se_text"), {**fmt, "bold": True, "italic": True})
            elif g("strong") is not None:
                inner = g("s_text") if g("s_text") is not None else g("s_text2")
                out += self._parse_inline(inner, {**fmt, "bold": True})
            elif g("em") is not None:
                inner = g("e_text") if g("e_text") is not None else g("e_text2")
                out += self._parse_inline(inner, {**fmt, "italic": True})
            elif g("strike") is not None:
                out += self._parse_inline(g("st_text"), {**fmt, "strike": True})
        if pos < len(text):
            out.append(("text", text[pos:], fmt))
        return out

    def _inline(self, p, text, size=None, bold=False, italic=False, color=None,
                muted=False, code_shrink=True):
        """Add `text` (inline Markdown) to paragraph `p` as formatted runs."""
        if not text:
            return
        base = {}
        if bold:
            base["bold"] = True
        if italic:
            base["italic"] = True
        self._emit(p, self._parse_inline(text, base), size=size, color=color,
                   muted=muted, code_shrink=code_shrink)

    def _emit(self, p, segments, size=None, color=None, muted=False, code_shrink=True):
        for seg in segments:
            if seg[0] == "image":
                _, alt, src, _fmt = seg
                if not self._add_picture(p.add_run(), src):
                    self._image_fallback(p, alt, src)
                continue
            _, text, fmt = seg
            if not text:
                continue
            run = p.add_run(text)
            if fmt.get("bold"):
                run.bold = True
            if fmt.get("italic"):
                run.italic = True
            if fmt.get("strike"):
                run.font.strike = True
            if size:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = color
            elif muted:
                run.font.color.rgb = PLACEHOLDER_COLOR
            if fmt.get("code"):
                _set_fonts(run._r.get_or_add_rPr(), self.o.code_font, self.cjk_font)
                run.font.color.rgb = CODE_COLOR
                if code_shrink:
                    run.font.size = Pt(max((size or self.o.font_size) - 1, 6))
            url = fmt.get("link")
            if url:
                run.font.color.rgb = LINK_COLOR
                run.font.underline = True
                if not url.startswith("#"):
                    self._wrap_hyperlink(p, run, url)

    @staticmethod
    def _wrap_hyperlink(p, run, url):
        r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        link.append(run._r)       # moves the run (it was the last child of p)
        p._p.append(link)


# ─────────────────────────────────────────────
# Public API
# ─────────────────────────────────────────────

def convert(md_text: str, options: Optional[Options] = None, base_dir: str | Path = "."):
    """Convert a Markdown string and return a python-docx Document.
    `base_dir` is used to resolve relative image paths."""
    return Converter(options, base_dir).convert(md_text)


def md_to_docx(md_path, docx_path=None, options: Optional[Options] = None, **kwargs):
    """Convert a Markdown file to .docx. Extra keyword arguments are passed to
    Options, e.g. md_to_docx("a.md", toc=True, cjk_font="宋体").
    Returns the output path."""
    md_path = Path(md_path)
    docx_path = Path(docx_path) if docx_path else md_path.with_suffix(".docx")
    if options is None:
        options = Options(**kwargs)
    text = md_path.read_text(encoding="utf-8-sig")  # -sig: tolerate a BOM (Windows Notepad)
    doc = convert(text, options, base_dir=md_path.parent)
    doc.save(str(docx_path))
    return docx_path


# ─────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────

def _expand_inputs(patterns):
    """Expand wildcards ourselves so `*.md` also works in Windows cmd/PowerShell."""
    files = []
    for pat in patterns:
        if any(ch in pat for ch in "*?[") and not Path(pat).exists():
            matches = sorted(glob.glob(pat))
            files.extend(matches or [pat])
        else:
            files.append(pat)
    return files


def main(argv=None):
    parser = argparse.ArgumentParser(
        prog="md2docx",
        description="Convert Markdown files to Word (.docx), with proper Chinese/CJK support.",
        epilog="Examples:\n"
               "  md2docx report.md\n"
               "  md2docx report.md -o report.docx --toc\n"
               "  md2docx notes/*.md --cjk-font 宋体\n"
               "  md2docx thesis.md --template university-template.docx",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("inputs", nargs="+", metavar="INPUT",
                        help="Markdown file(s); wildcards allowed. "
                             "For a single file you may give the output path as a second argument.")
    parser.add_argument("-o", "--output", help="output .docx path (single input only)")
    parser.add_argument("--font", help=f"Latin font (default: {DEFAULT_FONT})")
    parser.add_argument("--cjk-font", help=f"Chinese/Japanese font (default: {DEFAULT_CJK_FONT})")
    parser.add_argument("--code-font", default=DEFAULT_CODE_FONT,
                        help=f"font for code (default: {DEFAULT_CODE_FONT})")
    parser.add_argument("--font-size", type=float, default=11,
                        help="body font size in pt (default: 11)")
    parser.add_argument("--toc", action="store_true", help="insert a table of contents at the top")
    parser.add_argument("--toc-title", default="Contents",
                        help='heading text for the table of contents (default: "Contents")')
    parser.add_argument("--hr-pagebreak", action="store_true",
                        help="turn --- into a page break instead of a horizontal line")
    parser.add_argument("--template", metavar="DOCX",
                        help="use an existing .docx as the style template (its content is discarded)")
    parser.add_argument("--version", action="version", version=f"%(prog)s {__version__}")
    args = parser.parse_args(argv)

    inputs = list(args.inputs)
    output = args.output
    # Backward compatible: `md2docx input.md output.docx`
    if not output and len(inputs) == 2 and inputs[1].lower().endswith(".docx"):
        output = inputs.pop()
    inputs = _expand_inputs(inputs)
    if output and len(inputs) > 1:
        parser.error("-o/--output can only be used with a single input file")
    if args.template and not Path(args.template).is_file():
        parser.error(f"template not found: {args.template}")

    options = Options(
        font=args.font, cjk_font=args.cjk_font, code_font=args.code_font,
        font_size=args.font_size, toc=args.toc, toc_title=args.toc_title,
        hr_pagebreak=args.hr_pagebreak, template=args.template,
    )

    failures = 0
    for item in inputs:
        src = Path(item)
        if not src.is_file():
            print(f"Error: file not found: {src}", file=sys.stderr)
            failures += 1
            continue
        dst = Path(output) if output else src.with_suffix(".docx")
        try:
            md_to_docx(src, dst, options)
            print(f"Converted: {src} -> {dst}")
        except PermissionError:
            print(f"Error: cannot write {dst} - is it open in Word? Close it and try again.",
                  file=sys.stderr)
            failures += 1
        except UnicodeDecodeError:
            print(f"Error: {src} is not UTF-8 encoded. Re-save it as UTF-8 and try again.",
                  file=sys.stderr)
            failures += 1
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
